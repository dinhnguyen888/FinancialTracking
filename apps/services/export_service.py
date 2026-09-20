import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import List, Optional
from pathlib import Path
from datetime import datetime
from apps.models.transaction import Income, Expense

class ExportService:
    @classmethod
    def export_to_excel(
        cls,
        incomes: List[Income],
        expenses: List[Expense],
        start_date: str,
        end_date: str,
        output_path: Optional[Path] = None
    ) -> Path:
        """
        Export Incomes, Expenses, and Financial Summary to Excel (.xlsx) matching WPF ClosedXML structure.
        """
        if output_path is None:
            filename = f"Bao_Cao_Tai_Chinh_{start_date}_den_{end_date}.xlsx"
            output_dir = Path.home() / "Downloads"
            if not output_dir.exists():
                output_dir = Path("./data")
                output_dir.mkdir(exist_ok=True)
            output_path = output_dir / filename

        wb = openpyxl.Workbook()
        
        # Styles
        header_font = Font(name="Arial", size=11, bold=True, color="000000")
        header_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1"),
        )
        title_font = Font(name="Arial", size=14, bold=True, color="1E293B")
        summary_header_font = Font(name="Arial", size=11, bold=True)

        # 1. Sheet Thu nhập
        ws_income = wb.active
        ws_income.title = "Thu nhập"
        ws_income.views.sheetView[0].showGridLines = True

        income_headers = ["Ngày", "Số tiền (VNĐ)", "Danh mục", "Nguồn thu", "Ghi chú"]
        ws_income.append(income_headers)
        for col in range(1, len(income_headers) + 1):
            cell = ws_income.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for inc in sorted(incomes, key=lambda x: x.date):
            ws_income.append([
                inc.date,
                inc.amount,
                inc.category.name if inc.category else "Chưa phân loại",
                inc.source,
                inc.description
            ])

        for row in range(2, ws_income.max_row + 1):
            ws_income.cell(row=row, column=2).number_format = "#,##0"

        cls._auto_fit_columns(ws_income)

        # 2. Sheet Chi tiêu
        ws_expense = wb.create_sheet(title="Chi tiêu")
        ws_expense.views.sheetView[0].showGridLines = True

        expense_headers = ["Ngày", "Số tiền (VNĐ)", "Danh mục", "Phương thức", "Thiết yếu", "Ghi chú"]
        ws_expense.append(expense_headers)
        for col in range(1, len(expense_headers) + 1):
            cell = ws_expense.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for exp in sorted(expenses, key=lambda x: x.date):
            ws_expense.append([
                exp.date,
                exp.amount,
                exp.category.name if exp.category else "Chưa phân loại",
                exp.payment_method,
                "Có" if exp.is_essential else "Không",
                exp.description
            ])

        for row in range(2, ws_expense.max_row + 1):
            ws_expense.cell(row=row, column=2).number_format = "#,##0"

        cls._auto_fit_columns(ws_expense)

        # 3. Sheet Tổng kết
        ws_summary = wb.create_sheet(title="Tổng kết")
        ws_summary.views.sheetView[0].showGridLines = True

        ws_summary.cell(row=1, column=1, value="Báo Cáo Tổng Kết Tài Chính").font = title_font
        ws_summary.cell(row=2, column=1, value=f"Kỳ báo cáo: {start_date} đến {end_date}")

        total_income = sum(i.amount for i in incomes)
        total_expense = sum(e.amount for e in expenses)
        balance = total_income - total_expense
        essential_expense = sum(e.amount for e in expenses if e.is_essential)
        non_essential_expense = sum(e.amount for e in expenses if not e.is_essential)

        summary_data = [
            ("Tổng thu nhập:", total_income),
            ("Tổng chi tiêu:", total_expense),
            ("Số dư:", balance),
            ("Chi tiêu thiết yếu:", essential_expense),
            ("Chi tiêu không thiết yếu:", non_essential_expense),
        ]

        start_row = 4
        for idx, (label, val) in enumerate(summary_data):
            curr_row = start_row + idx
            lbl_cell = ws_summary.cell(row=curr_row, column=1, value=label)
            val_cell = ws_summary.cell(row=curr_row, column=2, value=val)

            lbl_cell.font = summary_header_font
            lbl_cell.border = thin_border
            val_cell.font = summary_header_font
            val_cell.number_format = "#,##0"
            val_cell.border = thin_border

        cls._auto_fit_columns(ws_summary)

        wb.save(str(output_path))
        return output_path

    @staticmethod
    def _auto_fit_columns(ws):
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                val_str = str(cell.value or '')
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
